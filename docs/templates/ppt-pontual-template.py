"""
TEMPLATE OFICIAL — Apresentações Pontual Brasil Petróleo
=========================================================
Identidade visual aprovada 2026-07-24. Base: `Downloads/capa apresentação julho.pptx`.
Ver regra completa: docs/memoria/feedback-ppt-padrao-visual-pontual.md

USO:
    1. Copie este arquivo pro projeto.
    2. Edite APENAS as seções marcadas com # >>> CONTEUDO
    3. Não altere paleta, fontes, header, footer sem checar com user.

DEPENDÊNCIA: pip install python-pptx pywin32
"""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# =============================================================
# PALETA OFICIAL PONTUAL  (--accent do design system frontend)
# =============================================================
NAVY   = RGBColor(0x18, 0x21, 0x6E)   # accent-700 institucional
NAVY_D = RGBColor(0x0F, 0x14, 0x43)   # accent-900
ICE    = RGBColor(0xDD, 0xE3, 0xF7)   # accent-100
ICE2   = RGBColor(0xEE, 0xF1, 0xFB)   # accent-50
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAY   = RGBColor(0x4A, 0x4A, 0x4A)
GRAY_L = RGBColor(0x8A, 0x8A, 0x8A)
GOLD   = RGBColor(0xF5, 0xB8, 0x00)   # amarelo Pontual (destaque)
AZUL_M = RGBColor(0x5F, 0x80, 0xCD)

FONTE_TITULO = "Calibri"
FONTE_CORPO  = "Calibri"
FONTE_PLACA  = "Consolas"   # única exceção (placas de veículo)

# =============================================================
# HELPERS BASE — nunca alterar sem revisar identidade visual
# =============================================================
def novo_deck():
    """Cria Presentation widescreen 13.333 x 7.5 pol (padrão Pontual)."""
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def novo_slide(prs):
    """Slide em branco com fundo BRANCO."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    return s

def add_rect(slide, x, y, w, h, fill, line=None, corner=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if corner else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if corner is not None:
        try: s.adjustments[0] = corner
        except Exception: pass
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.5)
    s.shadow.inherit = False
    return s

def add_text(slide, x, y, w, h, text, size=14, bold=False, color=None,
             align=PP_ALIGN.LEFT, font=None, anchor=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    if anchor is not None: tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font or FONTE_CORPO
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color or GRAY
    return tb

def header(slide, titulo, subtitulo, sw):
    """Faixa navy + faixa amarela + título CAIXA ALTA + subtítulo ice."""
    add_rect(slide, 0, 0, sw, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), sw, Inches(0.06), GOLD)
    add_text(slide, Inches(0.5), Inches(0.22), sw-Inches(1), Inches(0.5),
             titulo.upper(), size=24, bold=True, color=WHITE, font=FONTE_TITULO)
    add_text(slide, Inches(0.5), Inches(0.68), sw-Inches(1), Inches(0.32),
             subtitulo, size=12, color=ICE, font=FONTE_TITULO)

def footer(slide, txt, sw, sh):
    add_text(slide, Inches(0.5), sh-Inches(0.32), sw-Inches(1), Inches(0.22),
             txt, size=9, color=GRAY_L)

def kpi_card_ice(slide, x, y, w, h, label, valor, sub=""):
    """Card azul-claro com métrica principal (topo do slide)."""
    add_rect(slide, x, y, w, h, ICE, corner=0.12)
    add_text(slide, x+Inches(0.25), y+Inches(0.2), w-Inches(0.5), Inches(0.3),
             label.upper(), size=10, bold=True, color=NAVY)
    add_text(slide, x+Inches(0.25), y+Inches(0.5), w-Inches(0.5), Inches(0.55),
             valor, size=22, bold=True, color=NAVY, font=FONTE_TITULO)
    if sub:
        add_text(slide, x+Inches(0.25), h+y-Inches(0.4), w-Inches(0.5), Inches(0.25),
                 sub, size=9, color=GRAY_L)

def kpi_card_navy(slide, x, y, w, h, valor, label):
    """Card navy com destaque (número em amarelo)."""
    add_rect(slide, x, y, w, h, NAVY, corner=0.12)
    add_text(slide, x+Inches(0.25), y+Inches(0.1), w-Inches(0.5), Inches(0.55),
             valor, size=26, bold=True, color=GOLD, font=FONTE_TITULO)
    add_text(slide, x+Inches(0.25), y+h-Inches(0.4), w-Inches(0.5), Inches(0.3),
             label, size=10, bold=True, color=WHITE)

def secao_titulo(slide, x, y, w, titulo, sub=""):
    """Cabeçalho de seção interna (não é header do slide)."""
    add_text(slide, x, y, w, Inches(0.35),
             titulo.upper(), size=13, bold=True, color=NAVY)
    if sub:
        add_text(slide, x, y+Inches(0.32), w, Inches(0.28),
                 sub, size=10, color=GRAY_L)

def ranking_row(slide, x, y, w, pos, nome, sub_txt, valor_texto, valor_bar, valor_max):
    """Linha de ranking padrão: [pos card] [nome+sub] [barra] [valor]."""
    # posição
    fill = GOLD if pos == 1 else NAVY
    txt_color = NAVY if pos == 1 else WHITE
    add_rect(slide, x, y, Inches(0.6), Inches(0.35), fill, corner=0.15)
    add_text(slide, x, y+Inches(0.03), Inches(0.6), Inches(0.3),
             f"{pos}º", size=13, bold=True, color=txt_color,
             align=PP_ALIGN.CENTER, font=FONTE_TITULO)
    # nome CAIXA ALTA
    name_x = x + Inches(0.8)
    add_text(slide, name_x, y-Inches(0.02), Inches(5.8), Inches(0.28),
             nome.upper()[:42], size=11, bold=True, color=NAVY)
    add_text(slide, name_x, y+Inches(0.19), Inches(5.8), Inches(0.22),
             sub_txt, size=9, color=GRAY_L)
    # barra
    bar_x = x + Inches(6.1)
    bar_wmax = Inches(4.5)
    add_rect(slide, bar_x, y+Inches(0.11), bar_wmax, Inches(0.15), ICE)
    bar_w = int(bar_wmax * (valor_bar / valor_max))
    fill_bar = GOLD if pos == 1 else NAVY
    add_rect(slide, bar_x, y+Inches(0.11), bar_w, Inches(0.15), fill_bar)
    # valor
    add_text(slide, x + w - Inches(1.55), y+Inches(0.06), Inches(1.55), Inches(0.28),
             valor_texto, size=11, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)

# =============================================================
# FORMATADORES
# =============================================================
def brl(v):
    """R$ 1.234,56"""
    return f"R$ {v:,.2f}".replace(",","X").replace(".",",").replace("X",".")

def brk(v):
    """R$ 1,2k para gráficos apertados."""
    if v >= 1000:
        return f"R$ {v/1000:.1f}k".replace(".",",")
    return brl(v)

def title_case_pt(s):
    """Title Case respeitando conectivos PT + 'x' minúsculo nas rotas."""
    s = s.title()
    s = re.sub(r'(?<=\S )\b(Do|Da|De|Dos|Das|E|X)\b',
               lambda m: m.group(0).lower(), s)
    return s

# =============================================================
# RENDER (Windows + PowerPoint via COM)
# =============================================================
def render_slides_png(pptx_path, outdir, width=1600, height=900):
    """Exporta cada slide como PNG (só Windows com PowerPoint instalado)."""
    import win32com.client
    os.makedirs(outdir, exist_ok=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)
    try:
        for i, slide in enumerate(pres.Slides, 1):
            slide.Export(os.path.join(outdir, f"slide-{i}.png"),
                         "PNG", width, height)
    finally:
        pres.Close(); app.Quit()


# =============================================================
# EXEMPLO MÍNIMO — copie e adapte
# =============================================================
if __name__ == "__main__":
    prs = novo_deck()
    SW, SH = prs.slide_width, prs.slide_height

    # >>> CONTEUDO — troque tudo abaixo pela sua apresentação
    s1 = novo_slide(prs)
    header(s1, "Título do Slide", "Subtítulo · Contexto · Período", SW)

    # exemplo: 3 KPI cards ICE
    y0 = Inches(1.4); h0 = Inches(1.35); gap = Inches(0.2)
    cw = (SW - Inches(1.0) - gap*2) / 3
    kpi_card_ice(s1, Inches(0.5),               y0, cw, h0, "MÉTRICA 1", "R$ 100.000,00", "sub")
    kpi_card_ice(s1, Inches(0.5)+cw+gap,        y0, cw, h0, "MÉTRICA 2", "500", "sub")
    kpi_card_ice(s1, Inches(0.5)+(cw+gap)*2,    y0, cw, h0, "MÉTRICA 3", "42%", "sub")

    # exemplo: card navy destaque
    kpi_card_navy(s1, Inches(0.5), Inches(3.0), Inches(4), Inches(1.0),
                  "10x", "DESTAQUE PRINCIPAL")

    footer(s1, "Fonte: sistema interno · Pontual Brasil Petróleo", SW, SH)

    out = "exemplo_pontual.pptx"
    prs.save(out)
    print(f"OK: {out}")
