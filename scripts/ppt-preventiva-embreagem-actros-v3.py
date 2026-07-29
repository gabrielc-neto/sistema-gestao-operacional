"""
Slide único — Preventiva Embreagem New Actros (FORMATO 2)
Foto grande à esquerda + 1 parágrafo curto à direita. Sem colunas/cards.
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "docs", "templates"))
from importlib import import_module
tpl = import_module("ppt-pontual-template")

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY   = tpl.NAVY
NAVY_D = tpl.NAVY_D
GOLD   = tpl.GOLD
WHITE  = tpl.WHITE
GRAY   = RGBColor(0x33, 0x36, 0x4A)
GRAY_L = RGBColor(0x7A, 0x7E, 0x92)

FOTO = r"C:\Users\Logistica01\Pictures\Screenshots\Screenshot 2026-07-29 082318.png"

# ---- deck ----
prs = tpl.novo_deck()
SW, SH = prs.slide_width, prs.slide_height

s = prs.slides.add_slide(prs.slide_layouts[6])
tpl.add_rect(s, 0, 0, SW, SH, WHITE)

# ---- header ----
tpl.add_rect(s, 0, 0, SW, Inches(1.28), NAVY)
tpl.add_rect(s, 0, 0, Inches(0.08), Inches(1.28), GOLD)
tpl.add_text(s, Inches(0.5), Inches(0.25), SW-Inches(1), Inches(0.55),
             "DESCRIÇÃO DE MANUTENÇÃO PREVENTIVA · EMBREAGEM NEW ACTROS",
             size=24, bold=True, color=WHITE, font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.5), Inches(0.82), SW-Inches(1), Inches(0.35),
             "Frota Pontual Brasil Petróleo · 29/07/2026",
             size=12, color=RGBColor(0xC5,0xCC,0xE0), font=tpl.FONTE_TITULO)

# ---- FOTO grande à esquerda ----
foto_h = Inches(5.4)
foto_w = int(foto_h * 0.555)   # ~3.0"
foto_x = Inches(0.7)
foto_y = Inches(1.7)
s.shapes.add_picture(FOTO, foto_x, foto_y, foto_w, foto_h)
tpl.add_rect(s, foto_x, foto_y, foto_w, foto_h,
             RGBColor(0xFF,0xFF,0xFF), line=NAVY)

# ---- PARÁGRAFO à direita ----
tx = foto_x + foto_w + Inches(0.7)
tw = SW - tx - Inches(0.7)
ty = Inches(2.1)

# título curto acima
tpl.add_rect(s, tx, ty, Inches(0.7), Emu(30000), GOLD)
tpl.add_text(s, tx, ty+Inches(0.12), tw, Inches(0.5),
             "SOBRE A PREVENTIVA", size=14, bold=True, color=NAVY,
             font=tpl.FONTE_TITULO)

# parágrafo (um único bloco de texto corrido)
paragrafo = (
    "A embreagem do Mercedes-Benz New Actros passa por manutenção "
    "preventiva a cada 250.000 a 400.000 km, ou antes conforme o "
    "desgaste do material de fricção do disco. O serviço consiste em "
    "desmontar o conjunto (disco, platô e rolamento hidráulico), "
    "inspecionar o volante do motor e os retentores, medir o desgaste "
    "e substituir todas as peças que já não garantam a próxima janela "
    "de uso. O objetivo é fazer a troca em ambiente controlado, "
    "evitando parada em rota, garantindo custo previsível e menor "
    "tempo de veículo parado."
)
tf_box = s.shapes.add_textbox(tx, ty+Inches(0.75), tw, Inches(4.2))
tf = tf_box.text_frame
tf.margin_left = tf.margin_right = Emu(0)
tf.margin_top = tf.margin_bottom = Emu(0)
tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
p.line_spacing = 1.35
r = p.add_run(); r.text = paragrafo
r.font.name = tpl.FONTE_CORPO
r.font.size = Pt(14); r.font.color.rgb = GRAY

# ---- footer ----
tpl.add_text(s, Inches(0.5), SH-Inches(0.4), SW-Inches(1), Inches(0.25),
             "Fonte: manual Mercedes-Benz New Actros · Procedimento OS Pontual",
             size=9, color=GRAY_L)

out = r"C:\Users\Logistica01\Downloads\Preventiva_Embreagem_NewActros_2026_v3.pptx"
prs.save(out)
print(f"OK: {out}")
