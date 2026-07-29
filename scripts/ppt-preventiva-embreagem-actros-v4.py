"""
Slide único — SÓ título + foto (nada mais).
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "docs", "templates"))
from importlib import import_module
tpl = import_module("ppt-pontual-template")

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

NAVY = tpl.NAVY
GOLD = tpl.GOLD
WHITE = tpl.WHITE

FOTO = r"C:\Users\Logistica01\Pictures\Screenshots\Screenshot 2026-07-29 084032.png"

prs = tpl.novo_deck()
SW, SH = prs.slide_width, prs.slide_height

s = prs.slides.add_slide(prs.slide_layouts[6])
tpl.add_rect(s, 0, 0, SW, SH, WHITE)

# ---- header (só o título) ----
tpl.add_rect(s, 0, 0, SW, Inches(1.15), NAVY)
tpl.add_rect(s, 0, 0, Inches(0.08), Inches(1.15), GOLD)
tpl.add_text(s, Inches(0.5), Inches(0.35), SW-Inches(1), Inches(0.55),
             "MANUTENÇÃO PREVENTIVA EMBREAGEM NEW ACTROS",
             size=26, bold=True, color=WHITE, font=tpl.FONTE_TITULO)

# ---- foto centralizada (nada mais) ----
foto_h = Inches(6.0)
foto_w = int(foto_h * (469/842))   # proporção original (retrato)
foto_x = (SW - foto_w) / 2
foto_y = Inches(1.35)
s.shapes.add_picture(FOTO, foto_x, foto_y, foto_w, foto_h)

out = r"C:\Users\Logistica01\Downloads\Preventiva_Embreagem_NewActros_2026_v4.pptx"
prs.save(out)
print(f"OK: {out}")
