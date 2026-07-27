"""Lê SENHAS PARA VOCE.xlsx e gera CREDENCIAIS.md (gitignored)."""
import os, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

try:
    import openpyxl
except ImportError:
    print('instalando openpyxl...')
    os.system('pip install openpyxl --quiet')
    import openpyxl

XLSX = r'C:\Users\Logistica01\projetos\logistica-ia\arquivo\documents-fiscal\SENHAS PARA VOCE.xlsx'
OUT = r'C:\Users\Logistica01\projetos\logistica-ia\arquivo\CREDENCIAIS.md'

if not os.path.isfile(XLSX):
    print('arquivo não encontrado:', XLSX)
    sys.exit(1)

wb = openpyxl.load_workbook(XLSX, data_only=True)
lines = [
    '# 🔐 CREDENCIAIS Pontual Logística',
    '',
    '> ⚠️ **CONFIDENCIAL** — este arquivo está em `arquivo/` (gitignored). Nunca commitar.',
    '> Fonte: `arquivo/documents-fiscal/SENHAS PARA VOCE.xlsx`',
    '',
    '---',
    '',
]
for sheet in wb.sheetnames:
    ws = wb[sheet]
    lines.append(f'\n## 📋 {sheet}\n')
    for row in ws.iter_rows(values_only=True):
        nonempty = [str(c) if c is not None else '' for c in row]
        if any(v.strip() for v in nonempty):
            # Filtra linhas visualmente úteis
            lines.append('- ' + ' | '.join(v for v in nonempty if v))
    lines.append('')

# Salva
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))

# Só mostra que foi criado + total de linhas
print(f'✅ CREDENCIAIS.md criado ({os.path.getsize(OUT)} bytes, {len(lines)} linhas)')
print(f'   Local: {OUT}')
print(f'   Abas: {", ".join(wb.sheetnames)}')
