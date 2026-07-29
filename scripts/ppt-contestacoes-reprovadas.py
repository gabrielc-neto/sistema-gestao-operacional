"""
Slide único — Contestações reprovadas por MDF-e em aberto
Print da resposta da praça de pedágio + destaque do prejuízo.
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "docs", "templates"))
from importlib import import_module
tpl = import_module("ppt-pontual-template")

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY   = tpl.NAVY
NAVY_D = tpl.NAVY_D
GOLD   = tpl.GOLD
WHITE  = tpl.WHITE
GRAY   = RGBColor(0x33, 0x36, 0x4A)
GRAY_L = RGBColor(0x7A, 0x7E, 0x92)
ICE    = RGBColor(0xEE, 0xF1, 0xFB)
RED    = RGBColor(0xC0, 0x39, 0x2B)

PRINT_PATH = r"C:\Users\Logistica01\Pictures\Screenshots\Screenshot 2026-07-29 092508.png"
PRINT_RATIO = 1.882   # W/H

# ---- 2 contestações reprovadas (extraídas do print) ----
V1 = 53.04
V2 = 60.44
TOTAL = V1 + V2   # R$ 113,48

def brl(v):
    return f"R$ {v:,.2f}".replace(",","X").replace(".",",").replace("X",".")

prs = tpl.novo_deck()
SW, SH = prs.slide_width, prs.slide_height

s = prs.slides.add_slide(prs.slide_layouts[6])
tpl.add_rect(s, 0, 0, SW, SH, WHITE)

# ---- header ----
tpl.add_rect(s, 0, 0, SW, Inches(1.15), NAVY)
tpl.add_rect(s, 0, 0, Inches(0.08), Inches(1.15), GOLD)
tpl.add_text(s, Inches(0.5), Inches(0.2), SW-Inches(1), Inches(0.5),
             "CONTESTAÇÕES REPROVADAS · MDF-e EM ABERTO",
             size=22, bold=True, color=WHITE, font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.5), Inches(0.72), SW-Inches(1), Inches(0.4),
             "Retorno da praça de pedágio · Julho / 2026 · Placa SFL4G35",
             size=12, color=RGBColor(0xC5,0xCC,0xE0), font=tpl.FONTE_TITULO)

# ---- print (protagonista, ocupa quase o slide inteiro) ----
avail_h = Inches(4.4)
img_h = avail_h
img_w = int(img_h * PRINT_RATIO)
max_w = SW - Inches(1.0)
if img_w > max_w:
    img_w = max_w
    img_h = int(img_w / PRINT_RATIO)
img_x = (SW - img_w) // 2
img_y = Inches(1.35)

s.shapes.add_picture(PRINT_PATH, img_x, img_y, img_w, img_h)
tpl.add_rect(s, img_x, img_y, img_w, img_h,
             RGBColor(0xFF,0xFF,0xFF), line=NAVY)

# ---- cards de resumo embaixo ----
cy = Inches(5.95); ch = Inches(1.1); gap = Inches(0.2)

# 2 cards ICE (uma pra cada contestação)
cwsm = Inches(3.9)
tpl.add_rect(s, Inches(0.5), cy, cwsm, ch, ICE, corner=0.05)
tpl.add_text(s, Inches(0.65), cy+Inches(0.12), cwsm-Inches(0.3), Inches(0.28),
             "CS52173795 · 14/07/2026", size=10, bold=True, color=NAVY,
             font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.65), cy+Inches(0.4), cwsm-Inches(0.3), Inches(0.4),
             brl(V1), size=20, bold=True, color=NAVY_D, font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.65), cy+Inches(0.78), cwsm-Inches(0.3), Inches(0.25),
             "Porto Amazonas · BR277 158,38 · SFL4G35",
             size=9, color=GRAY_L)

tpl.add_rect(s, Inches(0.5)+cwsm+gap, cy, cwsm, ch, ICE, corner=0.05)
tpl.add_text(s, Inches(0.5)+cwsm+gap+Inches(0.15), cy+Inches(0.12),
             cwsm-Inches(0.3), Inches(0.28),
             "CS52150220 · 04/07/2026", size=10, bold=True, color=NAVY,
             font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.5)+cwsm+gap+Inches(0.15), cy+Inches(0.4),
             cwsm-Inches(0.3), Inches(0.4),
             brl(V2), size=20, bold=True, color=NAVY_D, font=tpl.FONTE_TITULO)
tpl.add_text(s, Inches(0.5)+cwsm+gap+Inches(0.15), cy+Inches(0.78),
             cwsm-Inches(0.3), Inches(0.25),
             "Irati · BR277 249,74 · SFL4G35",
             size=9, color=GRAY_L)

# card VERMELHO destaque com total
tx = Inches(0.5)+(cwsm+gap)*2
tw = SW - tx - Inches(0.5)
tpl.add_rect(s, tx, cy, tw, ch, RED, corner=0.05)
tpl.add_text(s, tx+Inches(0.25), cy+Inches(0.12), tw-Inches(0.5), Inches(0.28),
             "PREJUÍZO TOTAL · 2 CONTESTAÇÕES REPROVADAS",
             size=10, bold=True, color=WHITE, font=tpl.FONTE_TITULO)
tpl.add_text(s, tx+Inches(0.25), cy+Inches(0.4), tw-Inches(0.5), Inches(0.4),
             brl(TOTAL), size=24, bold=True, color=WHITE, font=tpl.FONTE_TITULO)
tpl.add_text(s, tx+Inches(0.25), cy+Inches(0.78), tw-Inches(0.5), Inches(0.25),
             "motivo: MDF-e não encerrado após conclusão da viagem",
             size=9, color=RGBColor(0xFF,0xE0,0xE0))

# footer
tpl.add_text(s, Inches(0.5), SH-Inches(0.4), SW-Inches(1), Inches(0.25),
             "Fonte: retorno da praça de pedágio · Julho/2026 · Pontual Brasil Petróleo",
             size=9, color=GRAY_L)

out = r"C:\Users\Logistica01\Downloads\Contestacoes_Reprovadas_MDFe_Julho_2026.pptx"
prs.save(out)
print(f"OK: {out}")
